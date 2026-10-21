import argparse,hashlib,json,pathlib,re
import openpyxl

# Default moved source; a positional path can still override this read-only input.
DEFAULT_SOURCE = pathlib.Path(r"C:\Users\chlgu\OneDrive\바탕 화면\한오엔 최종프로젝트\DB\한화오션_거제사업장_시설DB (1).xlsx")
from shipyard_value_status import apply_value_status, VALUE_STATUS_CODEBOOK

def run(source,out,reports):
    digest=hashlib.sha256(source.read_bytes()).hexdigest();w=openpyxl.load_workbook(source,read_only=True,data_only=True)
    out.mkdir(parents=True,exist_ok=True);reports.mkdir(parents=True,exist_ok=True)
    def write(name,records,sheet):
        (out/name).write_text(json.dumps({'sourceWorkbook':source.name,'sourceSha256':digest,'sheet':sheet,'records':records},ensure_ascii=False,indent=2),encoding='utf-8')
    def clean(v):return None if v is None or str(v).strip() in ('','미공개','미상','UNKNOWN') else v
    facilities=[]
    mapping=['id','image','name','officialName','englishName','classification','type','description','mainProcess','specifications','cranes','equipment','notes','confidence','source','imageFile']
    for num,row in enumerate(w['시설목록'].iter_rows(min_row=2,values_only=True),2):
        if not isinstance(row[0],(int,float)):continue
        f={k:clean(v) for k,v in zip(mapping,row)};f['sourceRow']=num;f['facilityId']=f'F{int(f["id"]):02}'
        specs={}
        for line in str(f['specifications'] or '').splitlines():
            if ':' in line:k,v=line.split(':',1);specs[k.strip()]=v.strip()
        for field,key in [('lengthM','길이(m)'),('widthM','폭(m)'),('depthM','깊이(m)'),('areaM2','적치장면적(㎡)'),('quayLengthM','안벽길이(m)')]:
            value=specs.get(key);f[field]=float(value.replace(',','')) if value and re.fullmatch(r'[\d,.]+',value) else None
        f['dimensionBasis']='WORKBOOK_PUBLIC_DATA' if any(f[k] is not None for k in ['lengthM','widthM','depthM','areaM2','quayLengthM']) else 'UNKNOWN'
        f['locationBasis']='CURRENT_LAYOUT_RECONSTRUCTION';facilities.append(f)
    assert [f['id'] for f in facilities]==list(range(1,23))
    cranes=[]
    for num,r in enumerate(w['크레인'].iter_rows(min_row=2,values_only=True),2):
        if not isinstance(r[0],(int,float)):continue
        cranes.append(dict(facilityId=f'F{int(r[0]):02}',facilityName=r[1],craneType=r[2],capacityText=str(r[3]),capacityTon=r[3] if isinstance(r[3],(int,float)) else None,quantity=r[4] if isinstance(r[4],(int,float)) else None,source=r[5],sourceRow=num,quantityBasis='CONFLICTING_TABLE_VS_TEXT' if r[0]==1 and r[2]=='집 크레인' else 'WORKBOOK_PUBLIC_DATA' if isinstance(r[4],(int,float)) else 'UNKNOWN'))
    floating=[];supplement=[]
    for num,r in enumerate(w['플로팅도크'].iter_rows(min_row=2,values_only=True),2):
        if r[0] and str(r[0]).startswith('Royal'):
            floating.append(dict(name=r[0],lengthM=r[1],widthM=r[2],depthM=r[3],capacityText=r[4],source=r[5],confidence='B',sourceRow=num))
        elif any(v is not None for v in r):supplement.append({'sourceRow':num,'cells':list(r)})
    overview=[{'key':r[0],'value':str(r[1]) if r[1] is not None else None,'sourceRow':i} for i,r in enumerate(w['사업장개요'].iter_rows(min_row=2,values_only=True),2) if r[0]]
    sources=[dict(key=r[0],name=r[1],confidence=r[2],url=r[3],sourceRow=i) for i,r in enumerate(w['출처'].iter_rows(min_row=2,values_only=True),2) if r[0] and r[3]]
    apply_value_status(facilities,floating)
    (out/'value_status_codebook.json').write_text(json.dumps(VALUE_STATUS_CODEBOOK,ensure_ascii=False,indent=2),encoding='utf-8')
    write('facility_master.json',facilities,'시설목록');write('crane_master.json',cranes,'크레인');write('floating_dock_master.json',floating,'플로팅도크');write('shipyard_overview.json',overview,'사업장개요');write('source_registry.json',sources,'출처')
    write('floating_dock_supplement.json',supplement,'플로팅도크')
    lines=['# Shipyard DB 실제 Schema 및 정규화','',f'원본: `{source}`',f'SHA-256: {digest}','','| Sheet | 물리 행 | 열 |','|---|---:|---:|']
    lines += [f'| {s.title} | {s.max_row} | {s.max_column} |' for s in w]
    lines += ['',f'실제 레코드: 시설 {len(facilities)}, 크레인 {len(cranes)}, Royal 플로팅도크 {len(floating)}, 사업장개요 {len(overview)}, URL 출처 {len(sources)}.',
              '', '## 실제 열과 값 형식']
    for s in w:
        rows=list(s.values);lines += ['',f'### {s.title}',str(list(rows[0]))]
        for c,name in enumerate(rows[0]):
            vals=[r[c] for r in rows[1:] if r[c] is not None];lines.append(f'- {name}: {sorted(set(type(v).__name__ for v in vals))}, 빈칸 {len(rows)-1-len(vals)}')
    lines += ['', '## 주의사항', '- 원본 Excel 미수정. 그림 셀은 빈 값이고 이미지 파일명은 별도 열에 있다. Runtime은 XLSX를 읽지 않는다.',
              '- F01 집크레인 수량: 표 8기 / 본문 6기 충돌을 보존했다. 모델의 표 기반 수량은 확정 보유대수로 해석하지 않는다.',
              '- 골리앗 최대 권상높이 78m와 전체 높이는 서로 다른 항목이다. 높이/폭 충돌을 메타데이터에 보존하며 외형은 ESTIMATED_FROM_PUBLIC_REFERENCES.',
              '- F09 3600m는 안벽 총길이로 취급한다. Scene에서는 합계 길이가 맞는 재구성 경로이며 실제 해안선 형상이 아니다.',
              '- F14 Royal 3기만 본체로 생성한다. 아래 증설 공사/사고 메모는 별도 JSON에 보존하며 현존 네 번째 도크로 만들지 않는다.',
              '- F20의 40×40m는 최대 블록 크기이며 공장 외형 치수가 아니다. F08 26km는 배관 길이이며 건물 치수가 아니다.',
              '- 크레인 용량 범위 30~50t는 원문 범위로 보존하고 단일 capacityTon을 만들지 않는다. 미공개 대수는 null.',
              '- 출처간 충돌·비고·신뢰등급은 그대로 보존했다. 외부 URL 내용의 재검증은 이번 정규화 작업과 별개이다.']
    (reports/'ShipyardDB_SchemaReport.md').write_text('\n'.join(lines),encoding='utf-8')
    assert hashlib.sha256(source.read_bytes()).hexdigest()==digest
    print(json.dumps({'facilities':len(facilities),'cranes':len(cranes),'floatingDocks':len(floating),'overview':len(overview),'sources':len(sources),'sourceUnchanged':True},ensure_ascii=False))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('source',type=pathlib.Path,nargs='?',default=DEFAULT_SOURCE);p.add_argument('--out',type=pathlib.Path,required=True);p.add_argument('--reports',type=pathlib.Path,required=True);a=p.parse_args();run(a.source,a.out,a.reports)
